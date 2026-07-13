# File name: scalping_catalyst_engine.py
# Scores weekly scalping catalyst records and runs offline overlay backtests.
from datetime import datetime
import math
import os
import re
import sqlite3

from openpyxl import load_workbook


EQUITY_SHEETS = {
    "US equities DB": "US",
    "CH equities DB": "CH",
    "KR equities DB": "KR",
    "JP equities DB": "JP",
}

OBJECTIVE_WEIGHTS = {
    "is_7pct_up": 15,
    "is_52w_high": 20,
    "is_estimate_up": 25,
    "is_turnover_top10": 15,
    "is_52w_low": -20,
}

ANALYST_MULTIPLIER = 0.20
ANALYST_CAP = 10
FIXED_OBJ_WEIGHT = 0.70
FIXED_ANA_WEIGHT = 0.30
FIXED_THRESHOLD = 35
FIXED_HOLDING_DAYS = 5
OBJECTIVE_MAX_SCORE = sum(weight for weight in OBJECTIVE_WEIGHTS.values() if weight > 0)
DB_PATH = os.path.join(os.path.dirname(__file__), "data", "market_data.sqlite")
_PRICE_CACHE = None
_TICKER_RESOLVE_CACHE = {}


def _clean(value):
    if value is None:
        return ""
    if hasattr(value, "date"):
        try:
            return value.date().isoformat()
        except Exception:
            pass
    return str(value).replace("\xa0", " ").strip()


def _parse_date(value):
    text = _clean(value)
    if not text:
        return None
    if "T" in text:
        text = text.split("T", 1)[0]
    match = re.match(r"^(\d{4,5})[-/.](\d{1,2})[-/.](\d{1,2})$", text)
    if not match:
        return None
    year, month, day = match.groups()
    if len(year) == 5 and year.startswith("202"):
        year = "2026"
    try:
        return datetime(int(year), int(month), int(day)).date()
    except Exception:
        return None


def _to_float(value):
    text = _clean(value).replace("%", "")
    if not text:
        return None
    try:
        return float(text)
    except Exception:
        return None


def _is_yes(value):
    text = _clean(value).upper()
    return text in {"Y", "YES", "TRUE", "1"}


def _score_analyst(linkage, size):
    if linkage is None or size is None:
        return 0.0
    raw = float(linkage) * abs(float(size)) / 100.0
    capped = min(raw * ANALYST_MULTIPLIER, ANALYST_CAP)
    if size < 0:
        return -capped
    return capped


def objective_score(record):
    score = 0.0
    for key, weight in OBJECTIVE_WEIGHTS.items():
        if record.get(key):
            score += weight
    return score


def analyst_raw_score(record):
    linkage = record.get("fundamental_linkage_score")
    size = record.get("catalyst_size_score")
    if linkage is None or size is None:
        return 0.0
    raw = float(linkage) * abs(float(size)) / 100.0
    return -raw if float(size) < 0 else raw


def capped_analyst_score(raw_score, analyst_cap=ANALYST_CAP):
    signed = -1.0 if raw_score < 0 else 1.0
    discounted = abs(raw_score) * ANALYST_MULTIPLIER
    return signed * min(discounted, analyst_cap)


def normalize_objective(score):
    if OBJECTIVE_MAX_SCORE <= 0:
        return 0.0
    return score / OBJECTIVE_MAX_SCORE * 100.0


def normalize_analyst(score, analyst_cap=ANALYST_CAP):
    if analyst_cap <= 0:
        return 0.0
    return score / analyst_cap * 100.0


def fixed_catalyst_score(record):
    obj_norm = normalize_objective(objective_score(record))
    ana = capped_analyst_score(analyst_raw_score(record), analyst_cap=ANALYST_CAP)
    ana_norm = normalize_analyst(ana, analyst_cap=ANALYST_CAP)
    return obj_norm * FIXED_OBJ_WEIGHT + ana_norm * FIXED_ANA_WEIGHT


def grade_score(score):
    if score >= 70:
        return "VERY_HIGH"
    if score >= 50:
        return "HIGH"
    if score >= 30:
        return "MEDIUM"
    if score >= 10:
        return "LOW"
    if score <= -30:
        return "NEGATIVE_HIGH"
    if score <= -10:
        return "NEGATIVE_WATCH"
    return "WEAK"


def score_record(record):
    objective = objective_score(record)

    analyst = _score_analyst(
        record.get("fundamental_linkage_score"),
        record.get("catalyst_size_score"),
    )
    total = objective + analyst
    return {
        "objective_score": objective,
        "analyst_score": analyst,
        "catalyst_score": total,
        "direction": "NEGATIVE" if total < 0 else "POSITIVE" if total > 0 else "NEUTRAL",
        "grade": grade_score(total),
    }


def _ticker_candidates(raw_ticker, country, market=None):
    ticker = _clean(raw_ticker).upper()
    if not ticker:
        return []
    ticker = ticker.replace(" HK EQUITY", "").replace(" CH EQUITY", "")
    ticker = ticker.replace(" JT EQUITY", "").replace(" KS EQUITY", "")
    ticker = ticker.split()[0]

    if "." in ticker:
        return [ticker]
    if country == "US":
        return [ticker]
    if country == "JP":
        return [f"{ticker}.T"]
    if country == "KR":
        base = ticker.zfill(6) if ticker.isdigit() else ticker
        return [f"{base}.KS", f"{base}.KQ"]
    if country == "CH":
        if market == "HK":
            base = ticker.zfill(4) if ticker.isdigit() else ticker
            return [f"{base}.HK"]
        if ticker.isdigit() and ticker.startswith("6"):
            return [f"{ticker.zfill(6)}.SS"]
        if ticker.isdigit():
            return [f"{ticker.zfill(6)}.SZ", f"{ticker.zfill(6)}.SS"]
    return [ticker]


def _load_price_cache():
    global _PRICE_CACHE
    if _PRICE_CACHE is not None:
        return _PRICE_CACHE
    _PRICE_CACHE = {}
    if not os.path.exists(DB_PATH):
        return _PRICE_CACHE

    conn = sqlite3.connect(DB_PATH)
    try:
        rows = conn.execute(
            """
            SELECT ticker, date, open, close
            FROM ohlcv
            ORDER BY ticker, date
            """
        ).fetchall()
    finally:
        conn.close()

    for ticker, date_value, open_value, close_value in rows:
        if open_value is None or close_value is None:
            continue
        try:
            _PRICE_CACHE.setdefault(ticker, []).append(
                (
                    datetime.strptime(str(date_value)[:10], "%Y-%m-%d").date(),
                    float(open_value),
                    float(close_value),
                )
            )
        except Exception:
            continue
    return _PRICE_CACHE


def _get_price_rows(ticker):
    if not ticker:
        return []
    return _load_price_cache().get(ticker, [])


def _resolve_cached_ticker(raw_ticker, country, market=None):
    cache_key = (raw_ticker, country, market)
    if cache_key in _TICKER_RESOLVE_CACHE:
        return _TICKER_RESOLVE_CACHE[cache_key]

    candidates = _ticker_candidates(raw_ticker, country, market=market)
    for candidate in candidates:
        if _get_price_rows(candidate):
            _TICKER_RESOLVE_CACHE[cache_key] = candidate
            return candidate
    resolved = candidates[0] if candidates else ""
    _TICKER_RESOLVE_CACHE[cache_key] = resolved
    return resolved


def load_scalping_events(workbook_path):
    wb = load_workbook(workbook_path, data_only=True, read_only=True)
    records = []

    for sheet_name, country in EQUITY_SHEETS.items():
        ws = wb[sheet_name]

        def value(row, idx, default=""):
            if idx >= len(row):
                return default
            return row[idx].value

        for row in ws.iter_rows(min_row=2):
            market_offset = 1 if country in {"CH", "KR"} else 0
            event_date = _parse_date(value(row, 0))
            market = _clean(value(row, 1)) if market_offset else ""
            raw_ticker = _clean(value(row, 1 + market_offset))
            if not event_date or not raw_ticker:
                continue

            record = {
                "event_date": event_date.isoformat(),
                "country": country,
                "market": market,
                "raw_ticker": raw_ticker,
                "ticker": _resolve_cached_ticker(raw_ticker, country, market=market),
                "company_name": _clean(value(row, 2 + market_offset)),
                "one_day_change_pct": _to_float(value(row, 3 + market_offset)),
                "is_7pct_up": _is_yes(value(row, 4 + market_offset)),
                "is_52w_high": _is_yes(value(row, 5 + market_offset)),
                "is_estimate_up": _is_yes(value(row, 6 + market_offset)),
                "is_turnover_top10": _is_yes(value(row, 7 + market_offset)),
                "is_52w_low": _is_yes(value(row, 8 + market_offset)),
                "fundamental_linkage_score": _to_float(value(row, 9 + market_offset)),
                "catalyst_size_score": _to_float(value(row, 10 + market_offset)),
                "author": _clean(value(row, 11 + market_offset)),
                "comment": _clean(value(row, 12 + market_offset)),
            }
            record.update(score_record(record))
            records.append(record)
    return records


def _forward_return(ticker, event_date, horizon_days):
    price_rows = _get_price_rows(ticker)
    if not price_rows:
        return None

    start_idx = None
    for idx, (date_value, _, _) in enumerate(price_rows):
        if date_value.isoformat() > event_date:
            start_idx = idx
            break
    if start_idx is None or start_idx + horizon_days >= len(price_rows):
        return None

    start_price = price_rows[start_idx][1]
    end_price = price_rows[start_idx + horizon_days][2]
    if start_price <= 0:
        return None
    return (end_price / start_price - 1.0) * 100.0


def _summary(values):
    values = [v for v in values if v is not None and not math.isnan(v)]
    if not values:
        return {"count": 0, "avg": 0.0, "std": 0.0, "win_rate": 0.0}
    wins = sum(1 for v in values if v > 0)
    avg = sum(values) / len(values)
    variance = sum((v - avg) ** 2 for v in values) / (len(values) - 1) if len(values) > 1 else 0.0
    return {
        "count": len(values),
        "avg": avg,
        "std": math.sqrt(variance),
        "win_rate": wins / len(values) * 100.0,
    }


def get_trade_returns(records, min_score=30, horizon_days=5, side="positive_long"):
    if side == "positive_long":
        selected = [r for r in records if r["catalyst_score"] >= min_score]
        multiplier = 1.0
    elif side == "negative_short":
        selected = [r for r in records if r["catalyst_score"] <= -min_score]
        multiplier = -1.0
    else:
        raise ValueError("side must be positive_long or negative_short")

    trades = []
    for record in selected:
        ret = _forward_return(record["ticker"], record["event_date"], horizon_days)
        if ret is None:
            continue
        trades.append(
            {
                "event_date": record["event_date"],
                "ticker": record["ticker"],
                "raw_ticker": record["raw_ticker"],
                "country": record["country"],
                "catalyst_score": record["catalyst_score"],
                "return_pct": ret * multiplier,
                "side": side,
                "horizon": f"{horizon_days}d",
            }
        )
    return trades


def get_fixed_trade_returns(records, horizon_days=FIXED_HOLDING_DAYS):
    trades = []
    for record in records:
        score = fixed_catalyst_score(record)
        if score < FIXED_THRESHOLD:
            continue
        ret = _forward_return(record["ticker"], record["event_date"], horizon_days)
        if ret is None:
            continue
        trades.append(
            {
                "event_date": record["event_date"],
                "ticker": record["ticker"],
                "raw_ticker": record["raw_ticker"],
                "country": record["country"],
                "fixed_score": score,
                "return_pct": ret,
                "side": "positive_long",
                "horizon": f"{horizon_days}d",
            }
        )
    return trades


def summarize_trades(trades):
    return _summary([trade["return_pct"] for trade in trades])


def backtest_catalyst(records, min_score=30, horizons=(1, 5)):
    positive = [r for r in records if r["catalyst_score"] >= min_score]
    negative = [r for r in records if r["catalyst_score"] <= -min_score]
    result = {"min_score": min_score, "positive": {}, "negative_short": {}}

    for horizon in horizons:
        pos_returns = [
            _forward_return(r["ticker"], r["event_date"], horizon) for r in positive
        ]
        neg_short_returns = []
        for record in negative:
            ret = _forward_return(record["ticker"], record["event_date"], horizon)
            neg_short_returns.append(-ret if ret is not None else None)
        result["positive"][f"{horizon}d"] = _summary(pos_returns)
        result["negative_short"][f"{horizon}d"] = _summary(neg_short_returns)
    return result
